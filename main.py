from fastapi import FastAPI, HTTPException, Depends, File, UploadFile
from pydantic import BaseModel
from typing import List, Optional
import uuid, os

app = FastAPI(title="AI-Assisted Doc Platform - Backend (scaffold)")

# ---- Simple in-memory 'DB' for scaffold/demo ----
DB = {"users":{}, "projects":{}}

class UserCreate(BaseModel):
    email: str
    password: str

class ProjectCreate(BaseModel):
    title: str
    doc_type: str  # "docx" or "pptx"
    outline: Optional[List[str]] = None  # for .docx
    slides: Optional[List[str]] = None   # for .pptx

@app.post("/register")
def register(user: UserCreate):
    uid = str(uuid.uuid4())
    DB["users"][uid] = {"email": user.email, "password": user.password}
    return {"user_id": uid, "message": "registered (scaffold, not secure)"}

@app.post("/projects")
def create_project(proj: ProjectCreate, user_id: str = "demo-user"):
    pid = str(uuid.uuid4())
    DB["projects"][pid] = {
        "title": proj.title,
        "doc_type": proj.doc_type,
        "outline": proj.outline or [],
        "slides": proj.slides or [],
        "content": {},  # section -> generated text
        "history": []
    }
    return {"project_id": pid, "project": DB["projects"][pid]}

@app.get("/projects/{project_id}")
def get_project(project_id: str):
    p = DB["projects"].get(project_id)
    if not p: raise HTTPException(404, "project not found")
    return p

# --- LLM stub endpoints (replace with real LLM calls) ---
@app.post("/generate/{project_id}/{section_index}")
def generate_section(project_id: str, section_index: int):
    p = DB["projects"].get(project_id)
    if not p: raise HTTPException(404, "project not found")
    key = f"section_{section_index}"
    text = f"Generated content for {p['title']} - section {section_index} (STUB)"
    p["content"][key] = text
    p["history"].append({"action":"generate","section":section_index,"text":text})
    return {"section": key, "text": text}

@app.post("/refine/{project_id}/{section_index}")
def refine_section(project_id: str, section_index: int, instruction: str):
    p = DB["projects"].get(project_id)
    if not p: raise HTTPException(404, "project not found")
    key = f"section_{section_index}"
    base = p["content"].get(key, f"(empty base for {key})")
    refined = f"Refined ({instruction}) -> {base}"
    p["content"][key] = refined
    p["history"].append({"action":"refine","section":section_index,"instruction":instruction,"text":refined})
    return {"section": key, "text": refined}

# --- Export endpoints (creates simple .docx/.pptx files) ---
from io import BytesIO
from fastapi.responses import StreamingResponse
from docx import Document
from pptx import Presentation

@app.get("/export/{project_id}")
def export_project(project_id: str):
    p = DB["projects"].get(project_id)
    if not p: raise HTTPException(404, "project not found")
    if p["doc_type"] == "docx":
        doc = Document()
        doc.add_heading(p["title"], level=1)
        # iterate outline sections
        for idx, sec in enumerate(p.get("outline", [])):
            doc.add_heading(sec, level=2)
            key = f"section_{idx}"
            doc.add_paragraph(p["content"].get(key, "(no generated content)"))
        bio = BytesIO()
        doc.save(bio)
        bio.seek(0)
        return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition":f"attachment; filename={p['title']}.docx"})
    else:
        prs = Presentation()
        # title slide
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = p["title"]
        subtitle.text = "Generated with scaffold"
        for idx, slide_title in enumerate(p.get("slides", [])):
            sl_layout = prs.slide_layouts[1]
            s = prs.slides.add_slide(sl_layout)
            s.shapes.title.text = slide_title
            key = f"section_{idx}"
            txBox = s.placeholders[1] if len(s.placeholders) > 1 else None
            content = p["content"].get(key, "(no content)")
            if txBox is not None:
                txBox.text = content
        bio = BytesIO()
        prs.save(bio)
        bio.seek(0)
        return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", headers={"Content-Disposition":f"attachment; filename={p['title']}.pptx"})